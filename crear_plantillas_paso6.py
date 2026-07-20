from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor


BASE = Path(__file__).resolve().parent
ORIGEN = BASE / "plantillas" / "plantilla_base_paso5.docx"
DESTINOS = (
    "plantilla_compraventa.docx",
    "plantilla_donacion.docx",
    "plantilla_cesion_derechos.docx",
)


def reemplazar_texto_parrafo(parrafo, texto: str) -> None:
    for run in list(parrafo.runs):
        parrafo._p.remove(run._r)
    run = parrafo.add_run(texto)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)


def crear() -> None:
    documento = Document(ORIGEN)

    p0 = documento.paragraphs[0]
    texto0 = p0.text
    inicio = texto0.find('Respecto de “LA TOTALIDAD”')
    fin = texto0.find(' - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - ANTECEDENTE')
    if inicio == -1 or fin == -1:
        raise RuntimeError("No se encontró la introducción del inmueble en la plantilla base.")
    texto0 = texto0[:inicio] + "{{ introduccion_acto }}" + texto0[fin:]
    reemplazar_texto_parrafo(p0, texto0)

    p1 = documento.paragraphs[1]
    texto1 = p1.text
    marca = "DATOS PERSONALES"
    indice = texto1.find(marca)
    if indice == -1:
        raise RuntimeError("No se encontró la sección DATOS PERSONALES.")
    cola = texto1[indice:]
    reemplazar_texto_parrafo(p1, "{{ clausulas_acto }} " + cola)

    p37 = documento.paragraphs[37]
    reemplazar_texto_parrafo(p37, "{{ cierre_apendice }}")

    for nombre in DESTINOS:
        documento.save(BASE / "plantillas" / nombre)


if __name__ == "__main__":
    crear()
