from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docxtpl import DocxTemplate, Listing
from docx.text.paragraph import Paragraph


class GeneradorWord:
    """Selecciona una plantilla por acto y genera el documento final."""

    def __init__(self, directorio_plantillas: str | Path) -> None:
        self.directorio_plantillas = Path(directorio_plantillas)
        if not self.directorio_plantillas.exists():
            raise FileNotFoundError(
                f"No se encontró el directorio de plantillas: {self.directorio_plantillas}"
            )

    def generar(
        self,
        datos: dict[str, Any],
        ruta_salida: str | Path,
        nombre_plantilla: str,
    ) -> Path:
        ruta_plantilla = self.directorio_plantillas / nombre_plantilla
        if not ruta_plantilla.exists():
            raise FileNotFoundError(f"No se encontró la plantilla: {ruta_plantilla}")

        ruta_salida = Path(ruta_salida)
        ruta_salida.parent.mkdir(parents=True, exist_ok=True)
        contexto = deepcopy(datos)

        self._convertir_a_listing(contexto, "inmueble", "medidas_colindancias")
        self._convertir_a_listing(contexto, None, "comparecientes_datos_personales")
        self._convertir_a_listing(contexto, None, "comparecientes_firmas")
        self._convertir_a_listing(contexto, None, "clausulas_acto")
        self._convertir_a_listing(contexto, None, "cierre_apendice")

        plantilla = DocxTemplate(ruta_plantilla)
        plantilla.render(contexto)
        plantilla.save(ruta_salida)

        nombres = [str(p.get("nombre", "")) for p in datos.get("comparecientes", [])]
        self._configurar_formato_final(ruta_salida, nombres)
        return ruta_salida

    @staticmethod
    def _convertir_a_listing(
        contexto: dict[str, Any],
        grupo: str | None,
        clave: str,
    ) -> None:
        contenedor = contexto if grupo is None else contexto.get(grupo, {})
        valor = contenedor.get(clave) if isinstance(contenedor, dict) else None
        if isinstance(valor, str):
            contenedor[clave] = Listing(valor)

    @staticmethod
    def _configurar_formato_final(
        ruta_documento: Path,
        nombres_comparecientes: list[str],
    ) -> None:
        documento = Document(ruta_documento)
        if not documento.sections:
            return

        seccion = documento.sections[0]
        ancho_util = seccion.page_width - seccion.left_margin - seccion.right_margin
        posicion_derecha = ancho_util - Pt(1)
        posicion_inicio = int(min(Inches(0.52), ancho_util // 5))
        posicion_centro = int(ancho_util // 2)

        nombres_limpios = {
            nombre.strip().upper() for nombre in nombres_comparecientes if nombre.strip()
        }

        # La plantilla original concentra varias secciones en párrafos enormes
        # separados con guiones escritos a mano. Se convierten en párrafos reales
        # antes de aplicar los líderes automáticos.
        GeneradorWord._separar_bloques_manual_guiones(documento)

        for parrafo in documento.paragraphs:
            texto_original = parrafo.text
            texto = GeneradorWord._limpiar_texto_visible(texto_original)
            if not texto:
                continue

            es_firma = GeneradorWord._es_parrafo_firma(
                texto_original,
                nombres_limpios,
            )

            if es_firma:
                texto_firma = re.sub(r"[ \n]+", " ", texto_original).strip()
                GeneradorWord._aplicar_formato_firma(
                    parrafo,
                    texto_firma,
                    posicion_derecha,
                )
                continue

            if texto.upper().startswith("ESCRITURA NÚMERO"):
                GeneradorWord._aplicar_encabezado_escritura(parrafo, texto)
            elif GeneradorWord._es_titulo_centrado(texto):
                GeneradorWord._aplicar_titulo_con_guiones(
                    parrafo,
                    texto,
                    posicion_centro,
                    posicion_derecha,
                )
            else:
                GeneradorWord._aplicar_parrafo_con_guiones(
                    parrafo,
                    texto,
                    posicion_inicio,
                    posicion_derecha,
                    ancho_util,
                )

            if re.match(r"^[a-zñ]\)\.-", texto.lower()):
                parrafo.paragraph_format.keep_together = True

        GeneradorWord._forzar_texto_negro(documento)
        GeneradorWord._eliminar_parrafos_vacios_finales(documento)
        documento.save(ruta_documento)

    @staticmethod
    def _separar_bloques_manual_guiones(documento: Any) -> None:
        patron = re.compile(r"(?:\s*-\s*){3,}")
        for parrafo in list(documento.paragraphs):
            texto = parrafo.text
            if not patron.search(texto):
                continue

            partes = [
                GeneradorWord._limpiar_texto_visible(parte)
                for parte in patron.split(texto)
            ]
            partes = [parte for parte in partes if parte]
            if len(partes) <= 1:
                continue

            padre = parrafo._p.getparent()
            posicion = padre.index(parrafo._p)
            ppr_original = parrafo._p.pPr

            for indice, parte in enumerate(partes):
                nuevo_p = OxmlElement("w:p")
                if ppr_original is not None:
                    nuevo_p.append(deepcopy(ppr_original))
                padre.insert(posicion + indice, nuevo_p)
                nuevo = Paragraph(nuevo_p, parrafo._parent)
                nuevo.add_run(parte)

            padre.remove(parrafo._p)

    @staticmethod
    def _limpiar_texto_visible(texto: str) -> str:
        texto = texto.replace("\t", " ").replace("\n", " ")
        texto = re.sub(r"(?:\s*-\s*){3,}", " ", texto)
        texto = re.sub(r"\s+", " ", texto)
        return texto.strip()

    @staticmethod
    def _es_titulo_centrado(texto: str) -> bool:
        normal = re.sub(r"\s+", " ", texto.strip().upper())
        compactado = normal.replace(" ", "")
        titulos = {
            "COMPRA VENTA",
            "DONACIÓN",
            "CESIÓN DE DERECHOS",
            "ANTECEDENTE",
            "DATOS PERSONALES",
            "CERTIFICACIÓN",
            "DOCUMENTOS DEL APÉNDICE",
        }
        compactos = {titulo.replace(" ", "") for titulo in titulos}
        if normal in titulos or compactado in compactos:
            return True
        if normal.startswith("LIBRO NÚMERO") and len(normal) < 120:
            return True
        return False


    @staticmethod
    def _es_parrafo_firma(
        texto_original: str,
        nombres_comparecientes: set[str],
    ) -> bool:
        texto = re.sub(r"[ \n]+", " ", texto_original).strip().upper()
        if not texto or not nombres_comparecientes:
            return False

        encontrados = [
            nombre for nombre in nombres_comparecientes if nombre and nombre in texto
        ]
        if not encontrados:
            return False

        resto = texto
        for nombre in sorted(encontrados, key=len, reverse=True):
            resto = resto.replace(nombre, "")
        resto = re.sub(r"[\s\t,;:.\-_/]+", "", resto)
        return not resto

    @staticmethod
    def _ancho_aproximado_puntos(texto: str) -> float:
        ancho = 0.0
        for caracter in texto:
            if caracter.isspace():
                ancho += 3.0
            elif caracter in "ilI.,;:'¡!|()[]{}":
                ancho += 3.1
            elif caracter in "MWÁÉÍÓÚÑmw@%":
                ancho += 8.0
            elif caracter.isupper() or caracter.isdigit():
                ancho += 6.5
            else:
                ancho += 5.5
        return ancho

    @staticmethod
    def _espacio_ultima_linea(
        texto: str,
        ancho_util: int,
        sangria_inicial: int,
    ) -> float:
        ancho_total = ancho_util / 12700.0
        ancho_primera = max(40.0, (ancho_util - sangria_inicial) / 12700.0)
        lineas: list[float] = []
        actual = 0.0
        limite = ancho_primera

        for palabra in texto.split():
            ancho_palabra = GeneradorWord._ancho_aproximado_puntos(palabra)
            separador = 0.0 if actual == 0 else 3.0
            if actual and actual + separador + ancho_palabra > limite:
                lineas.append(actual)
                actual = ancho_palabra
                limite = ancho_total
            else:
                actual += separador + ancho_palabra

        lineas.append(actual)
        limite_final = ancho_primera if len(lineas) == 1 else ancho_total
        return max(0.0, limite_final - lineas[-1])

    @staticmethod
    def _reemplazar_texto(parrafo: Any, texto: str) -> None:
        for run in list(parrafo.runs):
            parrafo._p.remove(run._r)
        run = parrafo.add_run(texto)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)

    @staticmethod
    def _aplicar_parrafo_con_guiones(
        parrafo: Any,
        texto: str,
        posicion_inicio: int,
        posicion_derecha: int,
        ancho_util: int,
    ) -> None:
        # Primer tabulador: franja corta de guiones a la izquierda.
        # Segundo tabulador: completa el espacio libre al margen derecho.
        # Si la última línea está prácticamente llena, se evita que Word mande
        # el líder a un renglón nuevo y se agregan solo los guiones que caben.
        restante = GeneradorWord._espacio_ultima_linea(
            texto, ancho_util, posicion_inicio
        )
        usar_lider_final = restante >= 100.0
        if usar_lider_final:
            texto_word = f"\t{texto}\t"
        else:
            cantidad = max(0, min(5, int(restante // 4.0)))
            texto_word = f"\t{texto}{'-' * cantidad}"

        GeneradorWord._reemplazar_texto(parrafo, texto_word)
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        parrafo.paragraph_format.keep_together = False
        tabs = parrafo.paragraph_format.tab_stops
        tabs.clear_all()
        tabs.add_tab_stop(
            posicion_inicio,
            WD_TAB_ALIGNMENT.LEFT,
            WD_TAB_LEADER.DASHES,
        )
        if usar_lider_final:
            tabs.add_tab_stop(
                posicion_derecha,
                WD_TAB_ALIGNMENT.RIGHT,
                WD_TAB_LEADER.DASHES,
            )
        GeneradorWord._uniformar_runs(parrafo)

    @staticmethod
    def _aplicar_encabezado_escritura(parrafo: Any, texto: str) -> None:
        # Este encabezado suele ser demasiado largo para dos líderes automáticos.
        # Se conserva en una sola línea, centrado y con guiones breves a ambos lados.
        tamano = 10.0 if len(texto) > 75 else 10.5
        GeneradorWord._reemplazar_texto(parrafo, f"--- {texto} ---")
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo.paragraph_format.keep_together = True
        parrafo.paragraph_format.tab_stops.clear_all()
        for run in parrafo.runs:
            run.font.name = "Arial"
            run.font.size = Pt(tamano)
            run.font.color.rgb = RGBColor(0, 0, 0)

    @staticmethod
    def _aplicar_titulo_con_guiones(
        parrafo: Any,
        texto: str,
        posicion_centro: int,
        posicion_derecha: int,
    ) -> None:
        GeneradorWord._reemplazar_texto(parrafo, f"\t{texto}\t")
        parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        parrafo.paragraph_format.keep_together = True
        tabs = parrafo.paragraph_format.tab_stops
        tabs.clear_all()
        tabs.add_tab_stop(
            posicion_centro,
            WD_TAB_ALIGNMENT.CENTER,
            WD_TAB_LEADER.DASHES,
        )
        tabs.add_tab_stop(
            posicion_derecha,
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DASHES,
        )
        GeneradorWord._uniformar_runs(parrafo)

    @staticmethod
    def _aplicar_formato_firma(
        parrafo: Any,
        texto: str,
        posicion_derecha: int,
    ) -> None:
        GeneradorWord._reemplazar_texto(parrafo, texto)
        parrafo.paragraph_format.keep_together = True
        tabs = parrafo.paragraph_format.tab_stops
        tabs.clear_all()
        if "\t" in texto:
            parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT
            tabs.add_tab_stop(
                posicion_derecha,
                WD_TAB_ALIGNMENT.RIGHT,
                WD_TAB_LEADER.SPACES,
            )
        else:
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        GeneradorWord._uniformar_runs(parrafo)


    @staticmethod
    def _eliminar_parrafos_vacios_finales(documento: Any) -> None:
        """Elimina párrafos vacíos al final que pueden crear una hoja en blanco."""
        body = documento.element.body
        elementos = list(body)
        for elemento in reversed(elementos):
            if elemento.tag == qn("w:sectPr"):
                continue
            if elemento.tag != qn("w:p"):
                break
            textos = elemento.xpath(".//w:t")
            contenido = "".join((nodo.text or "") for nodo in textos).strip()
            if contenido:
                break
            body.remove(elemento)

    @staticmethod
    def _forzar_texto_negro(documento: Any) -> None:
        for run_xml in documento.element.body.iter(qn("w:r")):
            rpr = run_xml.get_or_add_rPr()
            color = rpr.find(qn("w:color"))
            if color is None:
                color = OxmlElement("w:color")
                rpr.append(color)
            color.set(qn("w:val"), "000000")
            for atributo in (
                qn("w:themeColor"),
                qn("w:themeTint"),
                qn("w:themeShade"),
            ):
                if atributo in color.attrib:
                    del color.attrib[atributo]

    @staticmethod
    def _uniformar_runs(parrafo: Any) -> None:
        for run in parrafo.runs:
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)
